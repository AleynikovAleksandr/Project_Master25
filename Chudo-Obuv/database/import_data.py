import sys
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import create_engine, text

ROOT = Path(__file__).resolve().parent.parent
sys.path.append(str(ROOT))

import config

IMPORT_DIR = ROOT / "database" / "import"


def read_rows(file_name):
    sheet = load_workbook(IMPORT_DIR / file_name).active
    rows = list(sheet.iter_rows(values_only=True))
    return [row for row in rows[1:] if any(row)]


def clean(value):
    # В файлах заказчика есть лишние и неразрывные пробелы
    return " ".join(str(value).replace("\xa0", " ").split())


def get_id(conn, table, name):
    row = conn.execute(text(f"SELECT id FROM {table} WHERE name = :name"), {"name": name}).first()
    if row:
        return row.id
    return conn.execute(text(f"INSERT INTO {table} (name) VALUES (:name)"), {"name": name}).lastrowid


def find_product(products, name, manufacturer):
    # В заказах и остатках часть названий записана короче, чем в каталоге
    if (name, manufacturer) in products:
        return products[(name, manufacturer)]
    for (full_name, full_manufacturer), product_id in products.items():
        if full_manufacturer == manufacturer and full_name.startswith(name):
            return product_id
    raise ValueError(f"Товар не найден: {name} ({manufacturer})")


def create_tables(conn):
    sql = (ROOT / "database" / "schema.sql").read_text(encoding="utf-8")
    for statement in sql.split(";"):
        lines = [line for line in statement.splitlines() if not line.strip().startswith("--")]
        if "".join(lines).strip():
            conn.execute(text("\n".join(lines)))


def import_users(conn):
    for last_name, first_name, middle_name, login, role in read_rows("Users_import.xlsx"):
        conn.execute(
            text("INSERT INTO users (last_name, first_name, middle_name, login, role_id) "
                 "VALUES (:last_name, :first_name, :middle_name, :login, :role_id)"),
            {"last_name": clean(last_name), "first_name": clean(first_name),
             "middle_name": clean(middle_name), "login": clean(login),
             "role_id": get_id(conn, "roles", clean(role))},
        )


def import_products(conn):
    products = {}
    for category, subcategory, image, name, manufacturer, description, composition, price in read_rows("Products_import.xlsx"):
        product_id = conn.execute(
            text("INSERT INTO products (name, category_id, subcategory_id, manufacturer_id, "
                 "description, composition, price, image) VALUES (:name, :category_id, "
                 ":subcategory_id, :manufacturer_id, :description, :composition, :price, :image)"),
            {"name": clean(name), "category_id": get_id(conn, "categories", clean(category)),
             "subcategory_id": get_id(conn, "subcategories", clean(subcategory)),
             "manufacturer_id": get_id(conn, "manufacturers", clean(manufacturer)),
             "description": clean(description), "composition": clean(composition),
             "price": price, "image": clean(image) if image else None},
        ).lastrowid
        products[(clean(name), clean(manufacturer))] = product_id
    return products


def import_sizes(conn):
    sizes = {}
    for (size,) in read_rows("Sizes_import.xlsx"):
        sizes[clean(size)] = conn.execute(
            text("INSERT INTO sizes (value) VALUES (:value)"), {"value": clean(size)}
        ).lastrowid
    return sizes


def import_stock(conn, products, sizes):
    stock = {}
    for name, manufacturer, size, quantity in read_rows("Stock_Items_import.xlsx"):
        product_id = find_product(products, clean(name), clean(manufacturer))
        size_id = sizes[clean(size)]
        stock[(product_id, size_id)] = conn.execute(
            text("INSERT INTO stock_items (product_id, size_id, quantity) "
                 "VALUES (:product_id, :size_id, :quantity)"),
            {"product_id": product_id, "size_id": size_id, "quantity": quantity},
        ).lastrowid
    return stock


def import_orders(conn, products, sizes, stock):
    users = {}
    for row in conn.execute(text("SELECT id, last_name, first_name, middle_name FROM users")):
        users[f"{row.last_name} {row.first_name} {row.middle_name}"] = row.id

    created = set()
    for number, order_date, fio, category, name, manufacturer, size, quantity, price in read_rows("Orders_import.xlsx"):
        if clean(fio) not in users:
            raise ValueError(f"Клиент не найден среди пользователей: {fio}")
        if number not in created:
            created.add(number)
            conn.execute(
                text("INSERT INTO orders (id, order_date, user_id) VALUES (:id, :order_date, :user_id)"),
                {"id": number, "order_date": order_date.date(), "user_id": users[clean(fio)]},
            )
        product_id = find_product(products, clean(name), clean(manufacturer))
        conn.execute(
            text("INSERT INTO order_items (order_id, stock_item_id, quantity, unit_price) "
                 "VALUES (:order_id, :stock_item_id, :quantity, :unit_price)"),
            {"order_id": number, "stock_item_id": stock[(product_id, sizes[clean(size)])],
             "quantity": quantity, "unit_price": price},
        )


def main():
    engine = create_engine(config.DB_URL)
    with engine.begin() as conn:
        create_tables(conn)
        import_users(conn)
        products = import_products(conn)
        sizes = import_sizes(conn)
        stock = import_stock(conn, products, sizes)
        import_orders(conn, products, sizes, stock)
    print("Импорт завершён")


def prepare_database():
    server = create_engine(config.SERVER_URL)
    with server.begin() as conn:
        conn.execute(text(f"CREATE DATABASE IF NOT EXISTS `{config.DB_NAME}` CHARACTER SET utf8mb4"))

    engine = create_engine(config.DB_URL)
    with engine.connect() as conn:
        tables = conn.execute(text("SHOW TABLES")).all()
    if not tables:
        print(f"База {config.DB_NAME} пустая: создаю таблицы и загружаю данные из Excel")
        main()


if __name__ == "__main__":
    main()
